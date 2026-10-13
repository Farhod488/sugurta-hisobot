from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.core.models import SourceDocument, SourceTable


def new_document(title: str, source_name: str) -> Document:
    doc = Document()
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"Источник: {source_name}  ·  Сформировано: "
        f"{datetime.now().strftime('%d.%m.%Y %H:%M')}"
    )
    run.font.size = Pt(10)
    doc.add_paragraph()
    return doc


def add_section_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_key_value(doc: Document, key: str, value: object) -> None:
    paragraph = doc.add_paragraph()
    key_run = paragraph.add_run(f"{key}: ")
    key_run.bold = True
    paragraph.add_run(str(value))


def add_table(doc: Document, table: SourceTable, max_rows: int | None = None) -> None:
    rows = table.rows if max_rows is None else table.rows[:max_rows]
    if not table.columns:
        doc.add_paragraph("(пустая таблица)")
        return

    grid = doc.add_table(rows=1 + len(rows), cols=len(table.columns))
    grid.style = "Table Grid"

    for col_index, column in enumerate(table.columns):
        grid.rows[0].cells[col_index].text = str(column)

    for row_index, row in enumerate(rows, start=1):
        for col_index, column in enumerate(table.columns):
            grid.rows[row_index].cells[col_index].text = str(row.get(column, ""))


def save_document(doc: Document, output_dir: Path, filename: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / filename
    doc.save(target)
    return target


def apply_template_if_present(
    template_path: Path | None,
    title: str,
    source: SourceDocument,
) -> Document | None:
    """Load a DOCX template when provided; otherwise return None for default layout."""
    if template_path is None or not template_path.exists():
        return None
    doc = Document(template_path)
    # Keep template structure; append a generation stamp at the end.
    doc.add_paragraph()
    stamp = doc.add_paragraph()
    stamp.add_run(
        f"[Автозаполнение] Источник: {source.source_name}; "
        f"отчёт: {title}; "
        f"{datetime.now().strftime('%d.%m.%Y %H:%M')}"
    )
    return doc
