from __future__ import annotations

from pathlib import Path

from docx import Document

from app.core.models import SourceDocument, SourceTable


def read_docx(path: Path) -> SourceDocument:
    document = Document(path)
    text_blocks = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    tables: list[SourceTable] = []
    for index, table in enumerate(document.tables, start=1):
        if not table.rows:
            continue
        header_cells = table.rows[0].cells
        columns = [
            (cell.text.strip() or f"col_{i+1}")
            for i, cell in enumerate(header_cells)
        ]
        rows: list[dict[str, object]] = []
        for row in table.rows[1:]:
            cells = row.cells
            rows.append(
                {
                    columns[i]: (cells[i].text.strip() if i < len(cells) else "")
                    for i in range(len(columns))
                }
            )
        tables.append(
            SourceTable(name=f"table_{index}", columns=columns, rows=rows)
        )

    return SourceDocument(
        path=path,
        tables=tables,
        text_blocks=text_blocks,
        metadata={"format": "docx", "paragraph_count": len(text_blocks)},
    )
