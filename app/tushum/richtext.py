from __future__ import annotations

import re

from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

BLUE = RGBColor(0x00, 0x70, 0xC0)
RED = RGBColor(0xFF, 0x00, 0x00)
FONT_NAME = "Calibri"
FONT_SIZE = Pt(14)

_MONTHS = (
    "yanvar|fevral|mart|aprel|may|iyun|iyul|avgust|sentyabr|oktyabr|noyabr|dekabr"
)
_UNITS = r"ta|mln|ming|mlrd|soʻm|so'm|som"

_HIGHLIGHT_RE = re.compile(
    rf"(?P<number>\d+(?:[ \u00a0\u202f]\d{{3}})*(?:,\d+)?(?:-\d+(?:[ \u00a0\u202f]\d{{3}})*(?:,\d+)?)*)"
    rf"|(?P<unit>\b(?:{_UNITS})\b)"
    rf"|(?P<month>\b(?:{_MONTHS})\b)",
    re.IGNORECASE,
)


def _clear_runs(paragraph: Paragraph) -> None:
    for child in list(paragraph._element):
        if child.tag in {qn("w:r"), qn("w:hyperlink")}:
            paragraph._element.remove(child)


def _add_run(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    if text == "":
        return
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    rpr = run._element.get_or_add_rPr()
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = run._element.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), FONT_NAME)
    if color is not None:
        run.font.color.rgb = color


def _iter_segments(text: str):
    pos = 0
    for match in _HIGHLIGHT_RE.finditer(text):
        if match.start() > pos:
            yield "plain", text[pos : match.start()]
        kind = match.lastgroup or "plain"
        yield kind, match.group(0)
        pos = match.end()
    if pos < len(text):
        yield "plain", text[pos:]


def set_paragraph_rich(paragraph: Paragraph, text: str) -> None:
    """Rewrite paragraph using template number styles.

    - digits / dates: bold blue #0070C0
    - units (ta, mln, ming, mlrd, soʻm) and month names: bold red #FF0000
    - other text: normal Calibri 14
    """
    _clear_runs(paragraph)
    for kind, chunk in _iter_segments(text):
        if kind == "number":
            _add_run(paragraph, chunk, bold=True, color=BLUE)
        elif kind in {"unit", "month"}:
            _add_run(paragraph, chunk, bold=True, color=RED)
        else:
            _add_run(paragraph, chunk)
