from __future__ import annotations

import re
from dataclasses import dataclass, replace
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from app.tushum.assignments import (
    company_noun,
    filter_by_names,
    format_company_count,
    format_company_names_uz,
)
from app.tushum.formatters import (
    format_int_spaces,
    format_policy_count_words,
    mln_to_mlrd,
    soum_to_display,
    soum_to_mln,
)
from app.tushum.parsers import (
    ExcelTotals,
    PeriodTotals,
    filter_companies,
    parse_kuzatuv_pdf,
    parse_kuzatuv_xlsx,
    parse_osago_pdf,
    parse_osago_xlsx,
)
from app.tushum.profiles import PROFILES, Inspector, InspectorProfile, template_path
from app.tushum.richtext import set_paragraph_rich


# Content paragraph indices after MAʼLUMOT (same for personal and chief layouts)
IDX_KUZ_PDF = 3
IDX_KUZ_XLSX = 4
IDX_OSA_PDF = 7
IDX_OSA_XLSX = 8
IDX_SUM = 9


@dataclass
class TushumInputs:
    kuzatuv_pdf: Path | None
    kuzatuv_xlsx: Path | None
    osago_pdf: Path | None
    osago_xlsx: Path | None
    period_text: str  # e.g. "24-31-avgust"
    templates_dir: Path
    output_dir: Path
    # Exact Excel company names per inspector; None => profile defaults
    kuzatuv_assignments: dict[Inspector, list[str]] | None = None
    osago_assignments: dict[Inspector, list[str]] | None = None


def _optional_path(value: Path | str | None) -> Path | None:
    if value is None:
        return None
    path = Path(value)
    text = str(value).strip()
    if not text:
        return None
    return path


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    set_paragraph_rich(paragraph, text)


def _content_paragraphs(doc: Document) -> list[Paragraph]:
    """Return body paragraphs after MAʼLUMOT heading."""
    started = False
    result: list[Paragraph] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not started:
            normalized = text.upper().replace("ʼ", "'").replace("'", "")
            if "MALUMOT" in normalized:
                started = True
            continue
        if text:
            result.append(paragraph)
    return result


def _personal_kuz_pdf(kuz_total: PeriodTotals) -> str:
    return (
        f"Amaliyot yoʻlga qoʻyilganidan bugungi kungacha "
        f"{format_policy_count_words(kuz_total.policies)} ta "
        f"sugʻurta polislari asosida {mln_to_mlrd(kuz_total.premium_mln)} mlrd soʻm "
        f"sugʻurta mukofoti undirilgan. Shundan {mln_to_mlrd(kuz_total.mbl_share_mln)} "
        f"mlrd soʻm MBL rivojlantirish jamgʻarmasiga, "
        f"{mln_to_mlrd(kuz_total.company_share_mln)} mlrd soʻm sugʻurta tashkilotlari "
        f"hisob raqamlariga oʻtkazilgan."
    )


def _personal_osa_pdf(osa_total: PeriodTotals) -> str:
    return (
        f"Mazkur amaliyot yoʻlga qoʻyilishi natijasida, jami "
        f"{format_policy_count_words(osa_total.policies)} ta sugʻurta polislari "
        f"rasmiylashtirilib, {mln_to_mlrd(osa_total.premium_mln)} mlrd soʻm sugʻurta "
        f"mukofoti undirildi. Shundan {mln_to_mlrd(osa_total.mbl_share_mln)} mlrd soʻm "
        f"MBL, {mln_to_mlrd(osa_total.company_share_mln)} mlrd soʻm sugʻurta tashkiloti "
        f"hisobiga oʻtkazildi."
    )


def _aripov_kuz_xlsx(profile: InspectorProfile, period: str, kuz: ExcelTotals) -> str:
    prem_val, prem_unit = soum_to_display(kuz.premium_soum)
    count = profile.kuzatuv_company_count_text or ""
    names = profile.kuzatuv_company_names_text or ""
    noun = company_noun(int(count.split()[0]) if count else 0)
    return (
        f"Jumladan, joriy yilning {period} kunlari davomida Chet el tovarlarini "
        f"tashiyotgan avtotransport vositalarini bojxona hamrohligida kuzatib borishdan "
        f"ozod qilish uchun bojxona toʻlovlarini toʻlash majburiyatini sugʻurtalash "
        f"faoliyati yuzasidan bugungi kunda menga biriktirilgan "
        f"{count} {noun} ({names}) oʻrtasida imzolangan Texnik xizmat "
        f"koʻrsatish shartnomasi asosida jami {format_int_spaces(kuz.policies)} ta "
        f"sugʻurta polislari belgilangan tartibda rasmiylashtirilib, sugʻurta mukofoti "
        f"jami {prem_val} {prem_unit} soʻmni tashkil etdi. Shundan "
        f"{soum_to_mln(kuz.mbl_share_soum)} mln soʻm pul mablagʻlari Markaziy bojxona "
        f"laboratoriyasining rivojlantirish jamgʻarmasiga oʻtkazildi."
    )


def _aripov_osa_xlsx(profile: InspectorProfile, period: str, osa: ExcelTotals) -> str:
    prem_val, prem_unit = soum_to_display(osa.premium_soum)
    count = profile.osago_company_count_text or ""
    names = profile.osago_company_names_text or ""
    noun = company_noun(int(count.split()[0]) if count else 0)
    return (
        f"Jumladan, joriy yilning {period} kunlari davomida Chet el tovarlarini "
        f"tashiyotgan avtotransport vositalarini bojxona hamrohligida kuzatib borishdan "
        f"ozod qilish uchun bojxona toʻlovlarini toʻlash majburiyatini sugʻurtalash "
        f"faoliyati yuzasidan bugungi kunda menga biriktirilgan "
        f"{count} {noun} ({names}) oʻrtasida imzolangan Texnik xizmat "
        f"koʻrsatish shartnomasi asosida jami {format_int_spaces(osa.policies)} ta "
        f"sugʻurta polislari belgilangan tartibda rasmiylashtirilib, sugʻurta mukofoti "
        f"jami {prem_val} {prem_unit} soʻmni tashkil etdi. Shundan, Markaziy bojxona "
        f"laboratoriyasining rivojlantirish jamgʻarmasiga {soum_to_mln(osa.mbl_share_soum)} "
        f"mln soʻm undirildi."
    )


def _abdi_kuz_xlsx(profile: InspectorProfile, period: str, kuz: ExcelTotals) -> str:
    prem_val, prem_unit = soum_to_display(kuz.premium_soum)
    count = profile.kuzatuv_company_count_text or ""
    names = profile.kuzatuv_company_names_text or ""
    noun = company_noun(int(count.split()[0]) if count else 0)
    return (
        f"Jumladan, joriy yilning {period} kunlari davomida mazkur xizmat turi boʻyicha "
        f"bugungi kunda menga biriktirilgan {count} {noun} ({names}) oʻrtasida "
        f"imzolangan Texnik xizmat koʻrsatish shartnomasi asosida jami "
        f"{format_int_spaces(kuz.policies)} ta sugʻurta polislari belgilangan tartibda "
        f"rasmiylashtirilib, sugʻurta mukofoti jami {prem_val} {prem_unit} soʻmni "
        f"tashkil etdi. Shundan {soum_to_mln(kuz.mbl_share_soum)} mln soʻm pul "
        f"mablagʻlari Markaziy bojxona laboratoriyasining rivojlantirish jamgʻarmasiga "
        f"oʻtkazildi."
    )


def _abdi_osa_xlsx(profile: InspectorProfile, period: str, osa: ExcelTotals) -> str:
    prem_val, prem_unit = soum_to_display(osa.premium_soum)
    count = profile.osago_company_count_text or ""
    names = profile.osago_company_names_text or ""
    noun = company_noun(int(count.split()[0]) if count else 0)
    return (
        f"Jumladan, joriy yilning {period} kunlari davomida mazkur xizmat turi boʻyicha "
        f"bugungi kunda menga biriktirilgan {count} {noun} ({names}) oʻrtasida imzolangan "
        f"Texnik xizmat koʻrsatish shartnomasi asosida jami "
        f"{format_int_spaces(osa.policies)} ta sugʻurta polislari belgilangan tartibda "
        f"rasmiylashtirilib, sugʻurta mukofoti jami {prem_val} {prem_unit} soʻmni "
        f"tashkil etdi. Shundan, Markaziy bojxona laboratoriyasining rivojlantirish "
        f"jamgʻarmasiga {soum_to_mln(osa.mbl_share_soum)} mln soʻm undirildi."
    )


def _sum_paragraph(period: str, kuz: ExcelTotals, osa: ExcelTotals) -> str:
    total = soum_to_mln(kuz.mbl_share_soum + osa.mbl_share_soum)
    return (
        f"{period} kunlari davomida jami {total} mln soʻm pul mablagʻlari "
        f"Markaziy bojxona laboratoriyasining rivojlantirish jamgʻarmasiga oʻtkazildi."
    )


def _chief_kuz_xlsx(period: str, kuz: ExcelTotals) -> str:
    prem_val, prem_unit = soum_to_display(kuz.premium_soum)
    company_val, company_unit = soum_to_display(kuz.company_share_soum)
    return (
        f"Shuningdek, joriy yilning {period} kunlari mazkur xizmat turi boʻyicha "
        f"{format_int_spaces(kuz.policies)} ta polis asosida {prem_val} {prem_unit} "
        f"soʻm sugʻurta mukofoti undirilgan. Shundan {soum_to_mln(kuz.mbl_share_soum)} "
        f"mln soʻm MBL, {company_val} {company_unit} soʻm sugʻurta tashkilotlari hisob "
        f"raqamlariga oʻtkazilgan."
    )


def _chief_osa_xlsx(period: str, osa: ExcelTotals) -> str:
    company_val, company_unit = soum_to_display(osa.company_share_soum)
    return (
        f"Shuningdek, jori yilning {period} kunlari mazkur xizmat turi boʻyicha "
        f"{format_int_spaces(osa.policies)} ta polislar asosida "
        f"{soum_to_mln(osa.premium_soum)} mln soʻm sugʻurta mukofoti undirilgan. "
        f"Shundan {soum_to_mln(osa.mbl_share_soum)} mln soʻm MBL, {company_val} "
        f"{company_unit} soʻm sugʻurta tashkilotlari hisob raqamlariga oʻtkazilgan."
    )


def _build_updates(
    *,
    profile: InspectorProfile,
    period: str,
    kuz_total: PeriodTotals | None,
    osa_total: PeriodTotals | None,
    kuz_period: ExcelTotals | None,
    osa_period: ExcelTotals | None,
) -> dict[int, str]:
    updates: dict[int, str] = {}

    if kuz_total is not None:
        updates[IDX_KUZ_PDF] = _personal_kuz_pdf(kuz_total)

    if osa_total is not None:
        updates[IDX_OSA_PDF] = _personal_osa_pdf(osa_total)

    if profile.style == "personal":
        if kuz_period is not None:
            if profile.inspector == Inspector.ARIPOV:
                updates[IDX_KUZ_XLSX] = _aripov_kuz_xlsx(profile, period, kuz_period)
            else:
                updates[IDX_KUZ_XLSX] = _abdi_kuz_xlsx(profile, period, kuz_period)
        if osa_period is not None:
            if profile.inspector == Inspector.ARIPOV:
                updates[IDX_OSA_XLSX] = _aripov_osa_xlsx(profile, period, osa_period)
            else:
                updates[IDX_OSA_XLSX] = _abdi_osa_xlsx(profile, period, osa_period)
    else:
        # Chief: PDF paragraphs use the same wording as personal totals.
        if kuz_period is not None:
            updates[IDX_KUZ_XLSX] = _chief_kuz_xlsx(period, kuz_period)
        if osa_period is not None:
            updates[IDX_OSA_XLSX] = _chief_osa_xlsx(period, osa_period)

    if kuz_period is not None and osa_period is not None:
        updates[IDX_SUM] = _sum_paragraph(period, kuz_period, osa_period)

    return updates


def _replace_period_selective(
    doc: Document,
    period: str,
    *,
    protected_indices: set[int],
) -> None:
    """Replace sample period text, but never touch protected content paragraphs."""
    patterns = [
        r"24-31-avgust",
        r"24–31-avgust",
        r"24—31-avgust",
    ]
    content = _content_paragraphs(doc)
    protected_ids = {id(content[i]) for i in protected_indices if i < len(content)}

    for paragraph in doc.paragraphs:
        if id(paragraph) in protected_ids:
            continue
        original = paragraph.text
        if not original:
            continue
        updated = original
        for pattern in patterns:
            updated = re.sub(pattern, period, updated, flags=re.IGNORECASE)
        if updated != original:
            _set_paragraph_text(paragraph, updated)


def _assignment_names(
    assignments: dict[Inspector, list[str]] | None,
    inspector: Inspector,
) -> list[str] | None:
    if assignments is None or inspector not in assignments:
        return None
    return list(assignments[inspector])


def generate_tushum_report(
    inspector: Inspector | str,
    inputs: TushumInputs,
) -> Path:
    kind = Inspector(inspector) if isinstance(inspector, str) else inspector
    profile = PROFILES[kind]
    template = template_path(profile, inputs.templates_dir)
    if not template.exists():
        raise FileNotFoundError(f"Шаблон не найден: {template}")

    period = inputs.period_text.strip()
    if not period:
        raise ValueError("Укажите период (например: 24-31-avgust)")

    kuz_pdf = _optional_path(inputs.kuzatuv_pdf)
    osa_pdf = _optional_path(inputs.osago_pdf)
    kuz_xlsx = _optional_path(inputs.kuzatuv_xlsx)
    osa_xlsx = _optional_path(inputs.osago_xlsx)

    if not any([kuz_pdf, osa_pdf, kuz_xlsx, osa_xlsx]):
        raise ValueError(
            "Нужен хотя бы один файл: KUZATUV/OSAGO PDF или Excel. "
            "Без файла соответствующие абзацы останутся из шаблона."
        )

    for label, path in [
        ("KUZATUV PDF", kuz_pdf),
        ("OSAGO PDF", osa_pdf),
        ("KUZATUV Excel", kuz_xlsx),
        ("OSAGO Excel", osa_xlsx),
    ]:
        if path is not None and not path.exists():
            raise FileNotFoundError(f"{label} не найден: {path}")

    kuz_total = parse_kuzatuv_pdf(kuz_pdf) if kuz_pdf else None
    osa_total = parse_osago_pdf(osa_pdf) if osa_pdf else None

    kuz_names = _assignment_names(inputs.kuzatuv_assignments, kind)
    osa_names = _assignment_names(inputs.osago_assignments, kind)

    kuz_period = None
    if kuz_xlsx:
        loaded = parse_kuzatuv_xlsx(kuz_xlsx)
        if profile.style == "chief" and kuz_names is None:
            kuz_period = loaded
        elif kuz_names is not None:
            kuz_period = filter_by_names(
                loaded, kuz_names, source_label=kuz_xlsx.name
            )
            profile = replace(
                profile,
                kuzatuv_company_count_text=format_company_count(len(kuz_names)),
                kuzatuv_company_names_text=format_company_names_uz(kuz_names),
            )
        elif profile.kuzatuv_company_keywords:
            kuz_period = filter_companies(
                loaded,
                profile.kuzatuv_company_keywords,
                source_label=kuz_xlsx.name,
            )
        else:
            kuz_period = loaded

    osa_period = None
    if osa_xlsx:
        loaded = parse_osago_xlsx(osa_xlsx)
        if profile.style == "chief" and osa_names is None:
            osa_period = loaded
        elif osa_names is not None:
            osa_period = filter_by_names(
                loaded, osa_names, source_label=osa_xlsx.name
            )
            profile = replace(
                profile,
                osago_company_count_text=format_company_count(len(osa_names)),
                osago_company_names_text=format_company_names_uz(osa_names),
            )
        elif profile.osago_company_keywords:
            osa_period = filter_companies(
                loaded,
                profile.osago_company_keywords,
                source_label=osa_xlsx.name,
            )
        else:
            osa_period = loaded

    updates = _build_updates(
        profile=profile,
        period=period,
        kuz_total=kuz_total,
        osa_total=osa_total,
        kuz_period=kuz_period,
        osa_period=osa_period,
    )

    # Do not rewrite period inside paragraphs we intentionally leave untouched.
    protected: set[int] = set()
    if kuz_total is None:
        protected.add(IDX_KUZ_PDF)
    if osa_total is None:
        protected.add(IDX_OSA_PDF)
    if kuz_period is None:
        protected.add(IDX_KUZ_XLSX)
    if osa_period is None:
        protected.add(IDX_OSA_XLSX)
    if kuz_period is None or osa_period is None:
        protected.add(IDX_SUM)

    doc = Document(template)
    _replace_period_selective(doc, period, protected_indices=protected)
    content = _content_paragraphs(doc)
    for index, text in updates.items():
        if index >= len(content):
            raise ValueError(
                f"В шаблоне {template.name} недостаточно абзацев "
                f"(нужен индекс {index}, есть {len(content)})"
            )
        _set_paragraph_text(content[index], text)

    inputs.output_dir.mkdir(parents=True, exist_ok=True)
    safe_period = re.sub(r"[^\w\-]+", "_", period, flags=re.UNICODE)
    output = inputs.output_dir / f"{profile.output_prefix}_{safe_period}.docx"
    doc.save(output)
    return output
