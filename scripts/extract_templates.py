from pathlib import Path

import pandas as pd
import pdfplumber
from docx import Document

src = Path("templates/source")
out = Path("templates/_extracted")
out.mkdir(exist_ok=True)

for pdf in src.glob("*.pdf"):
    texts: list[str] = []
    with pdfplumber.open(pdf) as doc:
        for i, page in enumerate(doc.pages, 1):
            texts.append(f"=== PAGE {i} ===")
            texts.append(page.extract_text() or "")
            for ti, table in enumerate(page.extract_tables() or [], 1):
                texts.append(f"--- TABLE {ti} ---")
                for row in table:
                    cells = ["" if c is None else str(c).replace("\n", " ") for c in row]
                    texts.append(" | ".join(cells))
    (out / f"{pdf.stem}.txt").write_text("\n".join(texts), encoding="utf-8")
    print("pdf", pdf.name)

for xlsx in src.glob("*.xlsx"):
    xl = pd.ExcelFile(xlsx)
    chunks = [f"SHEETS: {xl.sheet_names}"]
    for sheet in xl.sheet_names:
        df = pd.read_excel(xlsx, sheet_name=sheet, header=None)
        chunks.append(f"\n=== SHEET {sheet} shape={df.shape} ===")
        # Limit huge sheets for inspection
        preview = df.head(80)
        chunks.append(preview.to_string())
        if len(df) > 80:
            chunks.append(f"... truncated, total rows={len(df)}")
    (out / f"{xlsx.stem}.xlsx.txt").write_text("\n".join(chunks), encoding="utf-8")
    print("xlsx", xlsx.name, "sheets", xl.sheet_names)

for docx_path in src.glob("*.docx"):
    if docx_path.name.startswith("~$"):
        continue
    d = Document(docx_path)
    lines: list[str] = []
    for i, p in enumerate(d.paragraphs, 1):
        style = p.style.name if p.style else ""
        text = p.text
        if text.strip():
            lines.append(f"P{i} [{style}]: {text}")
    for ti, table in enumerate(d.tables, 1):
        lines.append(f"\n--- DOCX TABLE {ti} ---")
        for row in table.rows:
            lines.append(" | ".join(c.text.replace("\n", " ") for c in row.cells))
    (out / f"{docx_path.stem}.docx.txt").write_text("\n".join(lines), encoding="utf-8")
    print("docx", docx_path.name)

print("done")
