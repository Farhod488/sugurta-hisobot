from pathlib import Path

import openpyxl
import pandas as pd

src = Path("templates/source")
out = Path("templates/_extracted/excel_detail.txt")
lines: list[str] = []

for name in ("KUZATUV.xlsx", "OSAGO.xlsx"):
    path = src / name
    lines.append("=" * 60)
    lines.append(f"{name} size={path.stat().st_size}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines.append(f"sheets: {wb.sheetnames}")
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(
            f"sheet={sheet_name!r} max_row={ws.max_row} max_col={ws.max_column}"
        )
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            lines.append(f"{i}: {row}")
            if i >= 40:
                lines.append("...truncated...")
                break
    wb.close()

    xl = pd.ExcelFile(path)
    for sheet in xl.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet, header=None)
        lines.append(f"pandas sheet={sheet!r} shape={df.shape}")
        lines.append(df.to_string())

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out, "lines", len(lines))
